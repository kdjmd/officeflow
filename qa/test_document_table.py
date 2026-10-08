# ruff: noqa: UP045 - tests mirror the engine's Python compatibility annotations.
from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from typing import Optional
from unittest import mock
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
ENGINE_PATH = ROOT / "app-source" / "document_table.py"
SPEC = importlib.util.spec_from_file_location("document_table", ENGINE_PATH)
assert SPEC and SPEC.loader
document_table = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = document_table
SPEC.loader.exec_module(document_table)


class DocumentTableTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.temp = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def write_text(self, name: str, content: str, encoding: str = "utf-8") -> Path:
        path = self.temp / name
        path.write_text(content, encoding=encoding)
        return path

    def write_docx_table(self, name: str, rows: list[list[str]]) -> Path:
        path = self.temp / name
        xml_rows = []
        for row in rows:
            xml_cells = "".join(
                f"<w:tc><w:p><w:r><w:t>{escape(value)}</w:t></w:r></w:p></w:tc>" for value in row
            )
            xml_rows.append(f"<w:tr>{xml_cells}</w:tr>")
        document_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body><w:tbl>{''.join(xml_rows)}</w:tbl><w:sectPr /></w:body></w:document>"
        )
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", document_xml)
        return path

    def analyze(self, files: list[Path], mode: str = "auto", fields: Optional[list] = None) -> dict:
        return document_table.analyze_request(
            {"files": [str(path) for path in files], "mode": mode, "fields": fields or []}
        )

    @staticmethod
    def table_by_kind(result: dict, kind: str) -> dict:
        return next(table for table in result["tables"] if table["kind"] == kind)

    def assert_exact_table(
        self, result: dict, kind: str, expected_columns: list[str], expected_rows: list[list[str]]
    ) -> dict:
        table = self.table_by_kind(result, kind)
        self.assertEqual([column["label"] for column in table["columns"]], expected_columns)
        actual_rows = [
            [row["cells"][column["key"]]["value"] for column in table["columns"]]
            for row in table["rows"]
        ]
        self.assertEqual(actual_rows, expected_rows)
        for row in table["rows"]:
            for cell in row["cells"].values():
                self.assertTrue(cell["sourceIds"])
                self.assertTrue(all(source_id in result["sources"] for source_id in cell["sourceIds"]))
        self.assertEqual(result["stats"]["sourceCoverage"], 1.0)
        self.assertFalse(result["stats"]["needsAi"])
        return table

    def test_generalization_inventory_markdown_uses_document_schema(self) -> None:
        source = self.write_text(
            "inventory.md",
            "|SKU|品名|数量|单价|\n|---|---|---|---|\n"
            "|P-001|炒锅|120|89.50|\n|P-002|煎锅|80|76.00|\n",
        )
        result = self.analyze([source])
        self.assert_exact_table(
            result,
            "markdown_table",
            ["SKU", "品名", "数量", "单价"],
            [["P-001", "炒锅", "120", "89.50"], ["P-002", "煎锅", "80", "76.00"]],
        )

    def test_generalization_contract_key_values_use_document_schema(self) -> None:
        source = self.write_text(
            "contracts.txt",
            "编号：HT-2026-001\n甲方：远山公司\n金额：128000元\n日期：2026-08-01\n\n"
            "编号：HT-2026-002\n甲方：海川公司\n金额：86000元\n日期：2026-08-15\n",
        )
        result = self.analyze([source])
        self.assert_exact_table(
            result,
            "key_value_records",
            ["编号", "甲方", "金额", "日期"],
            [
                ["HT-2026-001", "远山公司", "128000元", "2026-08-01"],
                ["HT-2026-002", "海川公司", "86000元", "2026-08-15"],
            ],
        )

    def test_generalization_project_tab_delimited_uses_document_schema(self) -> None:
        source = self.write_text(
            "projects.txt",
            "任务\t负责人\t状态\t截止\n"
            "完成原型\t周宁\t进行中\t2026-09-02\n"
            "上线验收\t沈清\t未开始\t2026-09-10\n",
        )
        result = self.analyze([source])
        self.assert_exact_table(
            result,
            "delimited_table",
            ["任务", "负责人", "状态", "截止"],
            [
                ["完成原型", "周宁", "进行中", "2026-09-02"],
                ["上线验收", "沈清", "未开始", "2026-09-10"],
            ],
        )

    def test_generalization_invoice_ocr_bbox_uses_dynamic_header(self) -> None:
        try:
            import fitz
        except ImportError:
            self.skipTest("PyMuPDF is unavailable")
        path = self.temp / "invoice-scan.pdf"
        pdf = fitz.open()
        pdf.new_page()
        pdf.save(path)
        pdf.close()

        def ocr_line(values: list[str], y: int) -> dict:
            x_positions = [60, 320, 600]
            return {
                "text": " ".join(values),
                "words": [
                    {
                        "text": value,
                        "bbox": {"x": x_positions[index], "y": y, "width": 150, "height": 50},
                    }
                    for index, value in enumerate(values)
                ],
            }

        payload = {
            "success": True,
            "language": "zh-Hans-CN",
            "imageWidth": 1000,
            "imageHeight": 500,
            "text": "票号 税额 价税合计\nFP001 82.64 1082.64\nFP002 45.28 545.28",
            "lines": [
                ocr_line(["票号", "税额", "价税合计"], 20),
                ocr_line(["FP001", "82.64", "1082.64"], 100),
                ocr_line(["FP002", "45.28", "545.28"], 180),
            ],
        }
        completed = subprocess.CompletedProcess(
            args=[], returncode=0, stdout=json.dumps(payload).encode("utf-8"), stderr=b""
        )
        with mock.patch.object(document_table.subprocess, "run", return_value=completed):
            result = self.analyze([path])
        self.assert_exact_table(
            result,
            "ocr_native_table",
            ["票号", "税额", "价税合计"],
            [["FP001", "82.64", "1082.64"], ["FP002", "45.28", "545.28"]],
        )

    def test_generalization_meeting_custom_fields_use_only_user_fields(self) -> None:
        source = self.write_text(
            "minutes.txt",
            "议题：发布排期\n决定：九月第一周发布\n行动项：完成回归测试\n\n"
            "议题：培训安排\n决定：采用线上培训\n行动项：整理培训材料\n",
        )
        fields = [
            {"key": "topic", "label": "议题"},
            {"key": "decision", "label": "决定"},
            {"key": "action", "label": "行动项"},
        ]
        result = self.analyze([source], mode="fields", fields=fields)
        self.assert_exact_table(
            result,
            "custom_fields",
            ["议题", "决定", "行动项"],
            [
                ["发布排期", "九月第一周发布", "完成回归测试"],
                ["培训安排", "采用线上培训", "整理培训材料"],
            ],
        )

    def test_generalization_log_docx_native_table_uses_document_schema(self) -> None:
        source = self.write_docx_table(
            "logs.docx",
            [
                ["时间", "级别", "事件"],
                ["10:00:01", "INFO", "服务启动"],
                ["10:05:12", "WARN", "重试一次"],
            ],
        )
        result = self.analyze([source])
        self.assert_exact_table(
            result,
            "docx_native_table",
            ["时间", "级别", "事件"],
            [["10:00:01", "INFO", "服务启动"], ["10:05:12", "WARN", "重试一次"]],
        )

    def test_markdown_table_is_local_and_traceable(self) -> None:
        source = self.write_text(
            "products.md",
            "|产品|规格|数量|\n|---|---|---|\n|炒锅|32cm|120|\n|煎锅|28cm|80|\n",
        )
        result = self.analyze([source])

        self.assertTrue(result["success"])
        table = self.table_by_kind(result, "markdown_table")
        self.assertEqual([column["label"] for column in table["columns"]], ["产品", "规格", "数量"])
        self.assertEqual(table["rows"][0]["cells"]["c1"]["value"], "炒锅")
        source_id = table["rows"][0]["cells"]["c1"]["sourceIds"][0]
        self.assertEqual(result["sources"][source_id]["lineStart"], 3)
        self.assertEqual(result["stats"]["sourceCoverage"], 1.0)
        self.assertFalse(result["stats"]["needsAi"])

    def test_key_value_records_and_prose_are_never_silently_dropped(self) -> None:
        source = self.write_text(
            "customers.txt",
            "客户：张三\n电话：13800000001\n金额：1280元\n\n"
            "客户：李四\n电话：13900000002\n金额：860元\n\n"
            "本月客户续约情况需要负责人进一步确认。\n",
        )
        result = self.analyze([source])

        key_value = self.table_by_kind(result, "key_value_records")
        ledger = self.table_by_kind(result, "paragraph_ledger")
        self.assertEqual(key_value["rowCount"], 2)
        self.assertEqual(ledger["rowCount"], 1)
        self.assertEqual(result["unresolved"][0]["text"], "本月客户续约情况需要负责人进一步确认。")
        self.assertEqual(result["stats"]["sourceCoverage"], 1.0)
        self.assertTrue(result["stats"]["needsAi"])

    def test_unrelated_key_value_schemas_are_split_into_separate_tables(self) -> None:
        source = self.write_text(
            "mixed.txt",
            "客户：张三\n电话：13800000001\n\n"
            "客户：李四\n电话：13900000002\n\n"
            "产品：炒锅\n数量：12\n\n"
            "产品：煎锅\n数量：8\n",
        )
        result = self.analyze([source])

        tables = [table for table in result["tables"] if table["kind"] == "key_value_records"]
        self.assertEqual(len(tables), 2)
        self.assertEqual([column["label"] for column in tables[0]["columns"]], ["客户", "电话"])
        self.assertEqual([column["label"] for column in tables[1]["columns"]], ["产品", "数量"])
        all_row_ids = [row["id"] for table in tables for row in table["rows"]]
        self.assertEqual(len(all_row_ids), len(set(all_row_ids)))

    def test_custom_fields_preserve_missing_values_and_typed_matches(self) -> None:
        source = self.write_text(
            "contacts.txt",
            "客户：张三\n电话：13800000001\n金额：1280元\n\n"
            "客户：李四\n电话：13900000002\n\n"
            "备用联系方式 contact@example.com，截止日期 2026-09-15。\n",
        )
        fields = [
            {"key": "name", "label": "客户", "aliases": ["姓名"]},
            {"key": "phone", "label": "电话"},
            {"key": "amount", "label": "金额"},
            {"key": "email", "label": "邮箱"},
            {"key": "date", "label": "截止日期"},
        ]
        result = self.analyze([source], mode="fields", fields=fields)

        kv_table = self.table_by_kind(result, "custom_fields")
        typed_table = self.table_by_kind(result, "custom_fields_text")
        self.assertEqual(kv_table["rowCount"], 2)
        self.assertTrue(kv_table["rows"][1]["cells"]["c3"]["missing"])
        self.assertEqual(typed_table["rows"][0]["cells"]["c4"]["value"], "contact@example.com")
        self.assertEqual(typed_table["rows"][0]["cells"]["c5"]["value"], "2026-09-15")
        self.assertGreaterEqual(result["stats"]["unresolvedCount"], 3)

    def test_docx_paragraph_and_native_table_without_word_dependency(self) -> None:
        path = self.temp / "sample.docx"
        document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>说明段落需要确认</w:t></w:r></w:p>
    <w:tbl>
      <w:tr><w:tc><w:p><w:r><w:t>名称</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>数量</w:t></w:r></w:p></w:tc></w:tr>
      <w:tr><w:tc><w:p><w:r><w:t>炒锅</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>12</w:t></w:r></w:p></w:tc></w:tr>
    </w:tbl>
    <w:sectPr />
  </w:body>
</w:document>"""
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", document_xml)

        result = self.analyze([path])

        native = self.table_by_kind(result, "docx_native_table")
        ledger = self.table_by_kind(result, "paragraph_ledger")
        self.assertEqual(native["rows"][0]["cells"]["c1"]["value"], "炒锅")
        self.assertEqual(ledger["rows"][0]["cells"]["c2"]["value"], "说明段落需要确认")
        source_id = native["rows"][0]["cells"]["c1"]["sourceIds"][0]
        self.assertEqual(result["sources"][source_id]["type"], "docx_table_cell")
        self.assertEqual(result["stats"]["sourceCoverage"], 1.0)

    def test_gb18030_text_is_decoded_without_data_loss(self) -> None:
        path = self.temp / "gb.txt"
        path.write_bytes("客户：王五\n电话：13800000003\n金额：560元\n".encode("gb18030"))
        result = self.analyze([path])
        table = self.table_by_kind(result, "key_value_records")
        self.assertEqual(table["rows"][0]["cells"]["c1"]["value"], "王五")
        self.assertEqual(result["stats"]["sourceCoverage"], 1.0)

    def test_text_pdf_is_extracted_and_analyzed(self) -> None:
        try:
            import fitz
        except ImportError:
            self.skipTest("PyMuPDF is unavailable")
        path = self.temp / "inventory.pdf"
        pdf = fitz.open()
        page = pdf.new_page()
        for y, line in ((72, "name,qty"), (92, "pan,12"), (112, "lid,8")):
            page.insert_text((72, y), line)
        pdf.save(path)
        pdf.close()

        result = self.analyze([path])

        self.assertTrue(result["success"], result["errors"])
        delimited = self.table_by_kind(result, "delimited_table")
        self.assertEqual(delimited["rowCount"], 2)
        self.assertEqual(delimited["rows"][1]["cells"]["c2"]["value"], "8")
        self.assertEqual(result["stats"]["sourceCoverage"], 1.0)

    def test_pdf_grid_table_is_preferred_as_native_table(self) -> None:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
        except ImportError:
            self.skipTest("reportlab is unavailable")
        path = self.temp / "grid.pdf"
        story_table = Table([["name", "qty"], ["pan", "12"], ["lid", "8"]])
        story_table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ]
            )
        )
        SimpleDocTemplate(str(path), pagesize=A4).build([story_table])

        result = self.analyze([path])
        native = self.table_by_kind(result, "pdf_native_table")
        self.assertEqual(native["rowCount"], 2)
        self.assertEqual(native["rows"][0]["cells"]["c1"]["value"], "pan")
        self.assertEqual(result["stats"]["sourceCoverage"], 1.0)

    def test_scanned_pdf_uses_local_ocr_and_continues_table_analysis(self) -> None:
        try:
            import fitz
        except ImportError:
            self.skipTest("PyMuPDF is unavailable")
        path = self.temp / "scan.pdf"
        pdf = fitz.open()
        pdf.new_page()
        pdf.save(path)
        pdf.close()

        ocr_payload = {
            "success": True,
            "language": "zh-Hans-CN",
            "imageWidth": 1200,
            "imageHeight": 600,
            "text": "姓名 部门 金额\n张三 市场部 1280\n李四 研发部 860",
            "lines": [
                {
                    "text": "姓 名 部 门 金 额",
                    "words": [
                        {"text": "姓", "bbox": {"x": 60, "y": 10, "width": 60, "height": 60}},
                        {"text": "名", "bbox": {"x": 125, "y": 10, "width": 55, "height": 60}},
                        {"text": "部", "bbox": {"x": 230, "y": 10, "width": 60, "height": 60}},
                        {"text": "门", "bbox": {"x": 295, "y": 10, "width": 55, "height": 60}},
                        {"text": "金", "bbox": {"x": 400, "y": 10, "width": 60, "height": 60}},
                        {"text": "额", "bbox": {"x": 465, "y": 10, "width": 55, "height": 60}},
                    ],
                },
                {
                    "text": "张 三 市 场 部 1280",
                    "words": [
                        {"text": "张", "bbox": {"x": 60, "y": 100, "width": 60, "height": 60}},
                        {"text": "三", "bbox": {"x": 125, "y": 100, "width": 55, "height": 60}},
                        {"text": "市", "bbox": {"x": 230, "y": 100, "width": 60, "height": 60}},
                        {"text": "场", "bbox": {"x": 295, "y": 100, "width": 60, "height": 60}},
                        {"text": "部", "bbox": {"x": 360, "y": 100, "width": 55, "height": 60}},
                        {"text": "1280", "bbox": {"x": 470, "y": 100, "width": 140, "height": 60}},
                    ],
                },
                {
                    "text": "李 四 研 发 部 860",
                    "words": [
                        {"text": "李", "bbox": {"x": 60, "y": 190, "width": 60, "height": 60}},
                        {"text": "四", "bbox": {"x": 125, "y": 190, "width": 55, "height": 60}},
                        {"text": "研", "bbox": {"x": 230, "y": 190, "width": 60, "height": 60}},
                        {"text": "发", "bbox": {"x": 295, "y": 190, "width": 60, "height": 60}},
                        {"text": "部", "bbox": {"x": 360, "y": 190, "width": 55, "height": 60}},
                        {"text": "860", "bbox": {"x": 470, "y": 190, "width": 110, "height": 60}},
                    ],
                },
            ],
        }
        completed = subprocess.CompletedProcess(
            args=[], returncode=0, stdout=json.dumps(ocr_payload).encode("utf-8"), stderr=b""
        )
        with mock.patch.object(document_table.subprocess, "run", return_value=completed) as mocked_run:
            result = self.analyze([path])

        self.assertTrue(result["success"])
        table = self.table_by_kind(result, "ocr_native_table")
        self.assertEqual(table["rowCount"], 2)
        self.assertEqual([column["label"] for column in table["columns"]], ["姓名", "部门", "金额"])
        self.assertEqual(table["rows"][0]["cells"]["c2"]["value"], "市场部")
        source_id = table["rows"][0]["cells"]["c1"]["sourceIds"][0]
        self.assertEqual(result["sources"][source_id]["type"], "ocr_table_cell")
        self.assertTrue(result["sources"][source_id]["ocr"])
        self.assertEqual(result["sources"][source_id]["bbox"]["x"], 60)
        self.assertFalse(result["stats"]["needsAi"])
        self.assertEqual(result["stats"]["sourceCoverage"], 1.0)
        command = mocked_run.call_args.args[0]
        temporary_png = Path(command[command.index("-InputFile") + 1])
        self.assertFalse(temporary_png.exists())

    def test_scanned_pdf_ocr_failure_is_explicitly_unresolved(self) -> None:
        try:
            import fitz
        except ImportError:
            self.skipTest("PyMuPDF is unavailable")
        path = self.temp / "scan-failed.pdf"
        pdf = fitz.open()
        pdf.new_page()
        pdf.save(path)
        pdf.close()

        payload = {
            "success": False,
            "error": {"code": "OCR_ENGINE_UNAVAILABLE", "message": "No language pack"},
        }
        completed = subprocess.CompletedProcess(
            args=[], returncode=3, stdout=json.dumps(payload).encode("utf-8"), stderr=b""
        )
        with mock.patch.object(document_table.subprocess, "run", return_value=completed):
            result = self.analyze([path])

        self.assertTrue(result["success"])
        self.assertEqual(result["unresolved"][0]["reason"], "OCR_REQUIRED")
        self.assertEqual(result["unresolved"][0]["suggestedAction"], "ocr")
        source_id = result["unresolved"][0]["sourceIds"][0]
        self.assertEqual(
            result["sources"][source_id]["ocrErrors"][0]["code"], "OCR_ENGINE_UNAVAILABLE"
        )
        self.assertTrue(result["stats"]["needsAi"])
        self.assertEqual(result["stats"]["sourceCoverage"], 1.0)

    def test_native_mode_lists_every_unhandled_source_as_unresolved(self) -> None:
        source = self.write_text("prose.txt", "第一段自由文本\n第二段自由文本\n")
        result = self.analyze([source], mode="native")

        self.assertEqual(result["tables"], [])
        self.assertEqual(len(result["unresolved"]), 2)
        self.assertEqual(result["stats"]["sourceCoverage"], 1.0)

    def test_xlsx_export_has_security_styles_and_traceability(self) -> None:
        source = self.write_text("formula.csv.txt", "name,value\nunsafe,=CMD()\nsafe,normal\n")
        analysis = self.analyze([source])
        target = self.temp / "result.xlsx"
        result = document_table.export_request(
            {"analysis": analysis, "outputPath": str(target), "format": "xlsx"}
        )

        from openpyxl import load_workbook

        workbook = load_workbook(target, data_only=False)
        self.assertTrue(result["success"])
        self.assertIn("数据", workbook.sheetnames)
        self.assertIn("待确认", workbook.sheetnames)
        self.assertIn("来源索引", workbook.sheetnames)
        data = workbook["数据"]
        self.assertEqual(data["B2"].value, "'=CMD()")
        self.assertEqual(data["B2"].data_type, "s")
        self.assertEqual(data.freeze_panes, "A2")
        self.assertIsNotNone(data["A2"].comment)
        self.assertGreater(workbook["来源索引"].max_row, 1)

    def test_a_table_title_cannot_displace_audit_sheet_names(self) -> None:
        analysis = {
            "tables": [
                {
                    "id": "table-0001",
                    "title": "主表",
                    "kind": "test",
                    "columns": [{"key": "c1", "label": "值"}],
                    "rows": [{"id": "r1", "cells": {"c1": {"value": "1", "sourceIds": []}}, "sourceIds": []}],
                },
                {
                    "id": "table-0002",
                    "title": "待确认",
                    "kind": "test",
                    "columns": [{"key": "c1", "label": "值"}],
                    "rows": [{"id": "r2", "cells": {"c1": {"value": "2", "sourceIds": []}}, "sourceIds": []}],
                },
            ],
            "unresolved": [],
            "warnings": [],
            "sources": {},
        }
        target = self.temp / "reserved.xlsx"
        document_table.export_request({"analysis": analysis, "outputPath": str(target), "format": "xlsx"})

        from openpyxl import load_workbook

        workbook = load_workbook(target)
        self.assertIn("待确认", workbook.sheetnames)
        self.assertIn("待确认_2", workbook.sheetnames)
        self.assertIn("来源索引", workbook.sheetnames)

    def test_csv_export_injects_no_formula(self) -> None:
        source = self.write_text("formula.txt", "name,value\nunsafe,+1+1\nsafe,normal\n")
        analysis = self.analyze([source])
        target = self.temp / "result.csv"
        document_table.export_request({"analysis": analysis, "outputPath": str(target), "format": "csv"})
        exported = target.read_text(encoding="utf-8-sig")
        self.assertIn("unsafe,'+1+1", exported)

    def test_cli_contract_writes_result_and_json_progress(self) -> None:
        source = self.write_text("simple.md", "|A|B|\n|---|---|\n|1|2|\n")
        request_path = self.temp / "request.json"
        result_path = self.temp / "result.json"
        request_path.write_text(
            json.dumps({"files": [str(source)], "mode": "auto", "fields": []}, ensure_ascii=False),
            encoding="utf-8",
        )
        completed = subprocess.run(
            [sys.executable, str(ENGINE_PATH), "analyze", str(request_path), str(result_path)],
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        progress = [json.loads(line) for line in completed.stdout.splitlines() if line.strip()]
        self.assertTrue(all(item["type"] == "progress" for item in progress))
        result = json.loads(result_path.read_text(encoding="utf-8"))
        self.assertTrue(result["success"])
        self.assertEqual(result["totalRows"], 1)
        self.assertEqual(result["rows"][0]["cells"]["c1"]["value"], "1")

    def test_invalid_extension_is_reported_in_errors(self) -> None:
        source = self.write_text("unsupported.rtf", "{\\rtf1 test}")
        result = self.analyze([source])
        self.assertFalse(result["success"])
        self.assertEqual(result["errors"][0]["code"], "UNSUPPORTED_FILE_TYPE")
        self.assertEqual(result["tables"], [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
