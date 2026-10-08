#!/usr/bin/env python3
# ruff: noqa: UP045 - Optional keeps the engine readable on older Python installs.
"""Local-first document-to-table analysis and export engine for OfficeFlow.

CLI contracts:
    python document_table.py analyze request.json result.json
    python document_table.py export request.json result.json

The engine deliberately keeps extraction deterministic.  Content that cannot be
mapped to a structured table is preserved in a ledger *and* listed in
``unresolved`` so callers can decide whether to ask a user or an AI model.
"""

from __future__ import annotations

import contextlib
import csv
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import time
import unicodedata
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional
from xml.etree import ElementTree

ENGINE_VERSION = "1.1.0"
SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}
MAX_REQUEST_BYTES = 5 * 1024 * 1024
MAX_INPUT_FILE_BYTES = 200 * 1024 * 1024
MAX_TEXT_FILE_BYTES = 50 * 1024 * 1024
MAX_FILES = 200
MAX_PDF_PAGES = 2_000
MAX_OCR_PAGES_PER_DOCUMENT = 50
OCR_TIMEOUT_SECONDS = 45
OCR_RENDER_SCALE = 2.0
MAX_DOCX_XML_BYTES = 64 * 1024 * 1024
MAX_CELL_CHARS = 32_767
FORMULA_PREFIXES = ("=", "+", "-", "@")
VALID_MODES = {"auto", "native", "ledger", "fields"}

WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{WORD_NS}}}"


class EngineError(Exception):
    """A safe, user-facing engine error."""

    def __init__(self, code: str, message: str, details: Optional[dict[str, Any]] = None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.details = details or {}

    def as_dict(self) -> dict[str, Any]:
        return {"code": self.code, "message": self.message, **self.details}


@dataclass
class TextUnit:
    text: str
    source: dict[str, Any]


@dataclass
class NativeTable:
    title: str
    rows: list[list[str]]
    cell_sources: list[list[dict[str, Any]]]
    kind: str = "native_table"
    has_header: Optional[bool] = True


@dataclass
class ExtractedDocument:
    path: str
    text_units: list[TextUnit] = field(default_factory=list)
    native_tables: list[NativeTable] = field(default_factory=list)
    warnings: list[dict[str, Any]] = field(default_factory=list)
    unresolved: list[dict[str, Any]] = field(default_factory=list)


class SourceRegistry:
    """Deduplicates source locations and assigns stable IDs."""

    def __init__(self) -> None:
        self._sources: dict[str, dict[str, Any]] = {}
        self._keys: dict[str, str] = {}

    def add(self, source: dict[str, Any]) -> str:
        clean = {key: value for key, value in source.items() if value is not None}
        key = json.dumps(clean, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        existing = self._keys.get(key)
        if existing:
            return existing
        source_id = f"src-{len(self._sources) + 1:06d}"
        clean["id"] = source_id
        self._sources[source_id] = clean
        self._keys[key] = source_id
        return source_id

    def as_dict(self) -> dict[str, dict[str, Any]]:
        return dict(self._sources)


def emit_progress(phase: str, current: int, total: int, message: str) -> None:
    payload = {
        "type": "progress",
        "phase": phase,
        "current": current,
        "total": total,
        "message": message,
    }
    encoded = (json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n").encode("utf-8")
    stream = getattr(sys.stdout, "buffer", None)
    if stream is not None:
        stream.write(encoded)
        stream.flush()
    else:
        sys.stdout.write(encoded.decode("utf-8"))
        sys.stdout.flush()


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    return text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n").strip()


def normalize_compare(value: str) -> str:
    return re.sub(r"\s+", "", normalize_text(value)).casefold()


def normalize_key(value: str) -> str:
    return re.sub(r"[\s:：_\-—（）()\[\]【】]+", "", normalize_text(value)).casefold()


def safe_cell_text(value: Any) -> str:
    text = normalize_text(value)
    if len(text) > MAX_CELL_CHARS:
        text = text[: MAX_CELL_CHARS - 1] + "…"
    stripped = text.lstrip(" \t\r\n\v\f")
    if stripped.startswith(FORMULA_PREFIXES):
        return "'" + text
    return text


def make_warning(code: str, message: str, **details: Any) -> dict[str, Any]:
    return {"code": code, "message": message, **{k: v for k, v in details.items() if v is not None}}


def base_source(path: str, source_type: str, text: str = "", **location: Any) -> dict[str, Any]:
    resolved = str(Path(path).resolve())
    return {
        "file": resolved,
        "fileName": Path(path).name,
        "type": source_type,
        **location,
        "text": normalize_text(text),
    }


def read_json_file(path: str) -> dict[str, Any]:
    request_path = Path(path)
    if not request_path.is_file():
        raise EngineError("REQUEST_NOT_FOUND", "请求文件不存在。", {"path": str(request_path)})
    if request_path.stat().st_size > MAX_REQUEST_BYTES:
        raise EngineError("REQUEST_TOO_LARGE", "请求 JSON 超过 5 MB 限制。")
    try:
        with request_path.open("r", encoding="utf-8-sig") as handle:
            value = json.load(handle)
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise EngineError("INVALID_JSON", "请求 JSON 无法解析。", {"detail": str(exc)}) from exc
    if not isinstance(value, dict):
        raise EngineError("INVALID_REQUEST", "请求 JSON 顶层必须是对象。")
    return value


def write_json_atomic(path: str, payload: dict[str, Any]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = ""
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            newline="\n",
            delete=False,
            dir=str(target.parent),
            prefix=f".{target.name}.",
            suffix=".tmp",
        ) as handle:
            temporary_path = handle.name
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        os.replace(temporary_path, target)
    except Exception:
        try:
            os.unlink(temporary_path)
        except OSError:
            pass
        raise


def validate_input_path(raw_path: Any) -> Path:
    if not isinstance(raw_path, str) or not raw_path.strip():
        raise EngineError("INVALID_FILE_PATH", "文件路径必须是非空字符串。")
    path = Path(raw_path).expanduser().resolve()
    if not path.is_file():
        raise EngineError("FILE_NOT_FOUND", "文件不存在或不是普通文件。", {"file": str(path)})
    if path.suffix.casefold() not in SUPPORTED_EXTENSIONS:
        raise EngineError(
            "UNSUPPORTED_FILE_TYPE",
            "仅支持 TXT、Markdown、DOCX 和文本型 PDF。",
            {"file": str(path), "extension": path.suffix.casefold()},
        )
    size = path.stat().st_size
    if size > MAX_INPUT_FILE_BYTES:
        raise EngineError(
            "FILE_TOO_LARGE",
            "单个文件超过 200 MB 限制。",
            {"file": str(path), "size": size},
        )
    if path.suffix.casefold() in {".txt", ".md"} and size > MAX_TEXT_FILE_BYTES:
        raise EngineError(
            "TEXT_FILE_TOO_LARGE",
            "TXT/Markdown 文件超过 50 MB 限制。",
            {"file": str(path), "size": size},
        )
    return path


def decode_plain_text(data: bytes) -> tuple[str, str]:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "utf-16"
    candidates = ("utf-8-sig", "gb18030")
    for encoding in candidates:
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8-replacement"


def extract_plain_text(path: Path) -> ExtractedDocument:
    document = ExtractedDocument(path=str(path))
    text, encoding = decode_plain_text(path.read_bytes())
    for line_number, line in enumerate(text.splitlines(), 1):
        document.text_units.append(
            TextUnit(line, base_source(str(path), "line", line, lineStart=line_number, lineEnd=line_number))
        )
    if "replacement" in encoding:
        document.warnings.append(
            make_warning("TEXT_ENCODING_REPLACED", "部分无法解码的字符已替换。", file=str(path))
        )
    if not any(unit.text.strip() for unit in document.text_units):
        document.warnings.append(make_warning("EMPTY_DOCUMENT", "文本文档没有可分析内容。", file=str(path)))
    return document


def word_text(element: ElementTree.Element) -> str:
    parts: list[str] = []
    for node in element.iter():
        if node.tag == W + "t":
            parts.append(node.text or "")
        elif node.tag == W + "tab":
            parts.append("\t")
        elif node.tag in {W + "br", W + "cr"}:
            parts.append("\n")
    return normalize_text("".join(parts))


def read_docx_xml(path: Path) -> bytes:
    try:
        with zipfile.ZipFile(path, "r") as archive:
            try:
                info = archive.getinfo("word/document.xml")
            except KeyError as exc:
                raise EngineError("INVALID_DOCX", "DOCX 缺少 word/document.xml。", {"file": str(path)}) from exc
            if info.file_size > MAX_DOCX_XML_BYTES:
                raise EngineError("DOCX_XML_TOO_LARGE", "DOCX 主文档 XML 超过 64 MB 限制。", {"file": str(path)})
            if info.compress_size and info.file_size / info.compress_size > 1_000:
                raise EngineError("DOCX_COMPRESSION_RATIO", "DOCX 压缩率异常，已拒绝处理。", {"file": str(path)})
            return archive.read(info)
    except zipfile.BadZipFile as exc:
        raise EngineError("INVALID_DOCX", "文件不是有效的 DOCX 压缩包。", {"file": str(path)}) from exc


def extract_docx(path: Path) -> ExtractedDocument:
    document = ExtractedDocument(path=str(path))
    xml_data = read_docx_xml(path)
    try:
        root = ElementTree.fromstring(xml_data)
    except ElementTree.ParseError as exc:
        raise EngineError("INVALID_DOCX_XML", "DOCX 主文档 XML 已损坏。", {"file": str(path)}) from exc
    body = root.find(".//" + W + "body")
    if body is None:
        raise EngineError("INVALID_DOCX_BODY", "DOCX 没有可读取的正文。", {"file": str(path)})

    paragraph_index = 0
    table_index = 0
    for child in list(body):
        if child.tag == W + "p":
            paragraph_index += 1
            text = word_text(child)
            document.text_units.append(
                TextUnit(
                    text,
                    base_source(str(path), "paragraph", text, paragraph=paragraph_index),
                )
            )
            continue
        if child.tag != W + "tbl":
            continue
        table_index += 1
        rows: list[list[str]] = []
        sources: list[list[dict[str, Any]]] = []
        for row_index, row_node in enumerate(child.findall("./" + W + "tr"), 1):
            row_values: list[str] = []
            row_sources: list[dict[str, Any]] = []
            for column_index, cell in enumerate(row_node.findall("./" + W + "tc"), 1):
                paragraphs = [word_text(node) for node in cell.findall("./" + W + "p")]
                text = normalize_text("\n".join(value for value in paragraphs if value))
                span = 1
                grid_span = cell.find("./" + W + "tcPr/" + W + "gridSpan")
                if grid_span is not None:
                    try:
                        span = max(1, min(100, int(grid_span.attrib.get(W + "val", "1"))))
                    except ValueError:
                        span = 1
                source = base_source(
                    str(path),
                    "docx_table_cell",
                    text,
                    table=table_index,
                    row=row_index,
                    column=column_index,
                )
                row_values.append(text)
                row_sources.append(source)
                for _ in range(span - 1):
                    row_values.append("")
                    row_sources.append({**source, "text": "", "mergedContinuation": True})
            if any(value for value in row_values):
                rows.append(row_values)
                sources.append(row_sources)
        if rows:
            document.native_tables.append(
                NativeTable(
                    title=f"{path.stem} - Word 表格 {table_index}",
                    rows=rows,
                    cell_sources=sources,
                    kind="docx_native_table",
                    has_header=True,
                )
            )

    if not document.native_tables and not any(unit.text for unit in document.text_units):
        document.warnings.append(make_warning("EMPTY_DOCUMENT", "DOCX 没有可分析内容。", file=str(path)))
    return document


def table_value_sets(tables: Sequence[NativeTable]) -> dict[int, set[str]]:
    values: dict[int, set[str]] = {}
    for table in tables:
        for source_row, row in zip(table.cell_sources, table.rows):
            page = int(source_row[0].get("page", 0)) if source_row else 0
            page_values = values.setdefault(page, set())
            normalized_cells = [normalize_compare(value) for value in row if normalize_compare(value)]
            page_values.update(normalized_cells)
            if normalized_cells:
                page_values.add("".join(normalized_cells))
    return values


def pdfplumber_tables(path: Path) -> tuple[list[NativeTable], Optional[str]]:
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        return [], "pdfplumber_unavailable"
    tables: list[NativeTable] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for page_number, page in enumerate(pdf.pages, 1):
                for table_number, raw_table in enumerate(page.extract_tables() or [], 1):
                    rows: list[list[str]] = []
                    sources: list[list[dict[str, Any]]] = []
                    for row_number, raw_row in enumerate(raw_table or [], 1):
                        if raw_row is None:
                            continue
                        row = [normalize_text(cell) for cell in raw_row]
                        if not any(row):
                            continue
                        rows.append(row)
                        sources.append(
                            [
                                base_source(
                                    str(path),
                                    "pdf_table_cell",
                                    value,
                                    page=page_number,
                                    table=table_number,
                                    row=row_number,
                                    column=column_number,
                                )
                                for column_number, value in enumerate(row, 1)
                            ]
                        )
                    if rows:
                        tables.append(
                            NativeTable(
                                title=f"{path.stem} - PDF 表格 P{page_number}-{table_number}",
                                rows=rows,
                                cell_sources=sources,
                                kind="pdf_native_table",
                                has_header=True,
                            )
                        )
    except Exception as exc:  # noqa: BLE001 - third-party parsers expose no stable exception base
        return [], f"pdfplumber_failed:{type(exc).__name__}"
    return tables, None


def pymupdf_tables(document: Any, path: Path) -> list[NativeTable]:
    tables: list[NativeTable] = []
    for page_number, page in enumerate(document, 1):
        find_tables = getattr(page, "find_tables", None)
        if not callable(find_tables):
            break
        try:
            # PyMuPDF may print an optional-package recommendation.  Suppress it
            # so stdout remains a strict JSON-lines progress channel.
            with contextlib.redirect_stdout(io.StringIO()):
                found = find_tables()
        except Exception:  # noqa: BLE001, S112 - a failed page falls back to text extraction
            continue
        for table_number, found_table in enumerate(getattr(found, "tables", []) or [], 1):
            try:
                raw_table = found_table.extract()
            except Exception:  # noqa: BLE001, S112 - preserve the rest of the document
                continue
            rows: list[list[str]] = []
            sources: list[list[dict[str, Any]]] = []
            for row_number, raw_row in enumerate(raw_table or [], 1):
                row = [normalize_text(cell) for cell in (raw_row or [])]
                if not any(row):
                    continue
                rows.append(row)
                sources.append(
                    [
                        base_source(
                            str(path),
                            "pdf_table_cell",
                            value,
                            page=page_number,
                            table=table_number,
                            row=row_number,
                            column=column_number,
                        )
                        for column_number, value in enumerate(row, 1)
                    ]
                )
            if rows:
                tables.append(
                    NativeTable(
                        title=f"{path.stem} - PDF 表格 P{page_number}-{table_number}",
                        rows=rows,
                        cell_sources=sources,
                        kind="pdf_native_table",
                        has_header=True,
                    )
                )
    return tables


def run_windows_ocr(png_path: Path, language: str = "zh-CN") -> dict[str, Any]:
    """Run the fixed local Windows OCR bridge and return its JSON payload."""
    script_path = Path(__file__).with_name("document_ocr.ps1")
    if not script_path.is_file():
        return {
            "success": False,
            "error": {"code": "OCR_SCRIPT_MISSING", "message": "本地 OCR 脚本不存在。"},
        }
    command = [
        "powershell.exe",
        "-NoLogo",
        "-NoProfile",
        "-NonInteractive",
        "-ExecutionPolicy",
        "Bypass",
        "-File",
        str(script_path),
        "-InputFile",
        str(png_path),
        "-Language",
        language,
    ]
    try:
        completed = subprocess.run(
            command,
            check=False,
            capture_output=True,
            timeout=OCR_TIMEOUT_SECONDS,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except subprocess.TimeoutExpired:
        return {
            "success": False,
            "error": {"code": "OCR_TIMEOUT", "message": f"本地 OCR 超过 {OCR_TIMEOUT_SECONDS} 秒。"},
        }
    except OSError as exc:
        return {
            "success": False,
            "error": {"code": "OCR_PROCESS_FAILED", "message": str(exc)},
        }

    stdout = completed.stdout.decode("utf-8-sig", errors="replace").strip()
    json_line = next((line for line in reversed(stdout.splitlines()) if line.lstrip().startswith("{")), "")
    if json_line:
        try:
            payload = json.loads(json_line)
            if isinstance(payload, dict):
                return payload
        except json.JSONDecodeError:
            pass
    stderr = completed.stderr.decode("utf-8-sig", errors="replace").strip()
    return {
        "success": False,
        "error": {
            "code": "OCR_INVALID_RESPONSE",
            "message": stderr[:500] or "本地 OCR 返回了无效结果。",
        },
    }


def numeric_box(value: Any) -> Optional[dict[str, float]]:
    if not isinstance(value, dict):
        return None
    try:
        return {
            "x": float(value.get("x", 0)),
            "y": float(value.get("y", 0)),
            "width": max(0.0, float(value.get("width", 0))),
            "height": max(0.0, float(value.get("height", 0))),
        }
    except (TypeError, ValueError):
        return None


def combine_ocr_tokens(tokens: Sequence[str]) -> str:
    text = ""
    for token in tokens:
        token = normalize_text(token)
        if not token:
            continue
        if not text:
            text = token
            continue
        previous = text[-1]
        current = token[0]
        previous_cjk = "\u3400" <= previous <= "\u9fff"
        current_cjk = "\u3400" <= current <= "\u9fff"
        separator = "" if previous_cjk or current_cjk else " "
        text += separator + token
    return text


def group_ocr_words(line: dict[str, Any]) -> list[dict[str, Any]]:
    positioned: list[tuple[str, dict[str, float]]] = []
    for word in line.get("words", []) if isinstance(line.get("words"), list) else []:
        if not isinstance(word, dict):
            continue
        text = normalize_text(word.get("text"))
        bbox = numeric_box(word.get("bbox"))
        if text and bbox:
            positioned.append((text, bbox))
    positioned.sort(key=lambda item: (item[1]["x"], item[1]["y"]))
    if len(positioned) < 2:
        return []
    heights = sorted(box["height"] for _text, box in positioned if box["height"] > 0)
    median_height = heights[len(heights) // 2] if heights else 20.0
    gap_threshold = max(8.0, min(40.0, median_height * 0.38))
    groups: list[list[tuple[str, dict[str, float]]]] = []
    current: list[tuple[str, dict[str, float]]] = []
    current_right = 0.0
    for text, bbox in positioned:
        gap = bbox["x"] - current_right if current else 0.0
        if current and gap > gap_threshold:
            groups.append(current)
            current = []
        current.append((text, bbox))
        current_right = max(current_right, bbox["x"] + bbox["width"]) if len(current) > 1 else bbox["x"] + bbox["width"]
    if current:
        groups.append(current)
    cells: list[dict[str, Any]] = []
    for group in groups:
        left = min(box["x"] for _text, box in group)
        top = min(box["y"] for _text, box in group)
        right = max(box["x"] + box["width"] for _text, box in group)
        bottom = max(box["y"] + box["height"] for _text, box in group)
        cells.append(
            {
                "text": combine_ocr_tokens([text for text, _box in group]),
                "bbox": {
                    "x": round(left, 2),
                    "y": round(top, 2),
                    "width": round(right - left, 2),
                    "height": round(bottom - top, 2),
                },
            }
        )
    return cells if len(cells) >= 2 else []


def infer_ocr_tables(
    lines: Sequence[dict[str, Any]],
    path: Path,
    page_number: int,
    language: str,
    image_width: float,
) -> tuple[list[NativeTable], set[int]]:
    """Infer simple scanned tables from repeated, aligned OCR word groups."""
    grouped = [(index, group_ocr_words(line)) for index, line in enumerate(lines)]
    tables: list[NativeTable] = []
    consumed: set[int] = set()
    position = 0
    while position < len(grouped):
        line_index, cells = grouped[position]
        if len(cells) < 2:
            position += 1
            continue
        region = [(line_index, cells)]
        next_position = position + 1
        while next_position < len(grouped):
            next_index, next_cells = grouped[next_position]
            if next_index != region[-1][0] + 1 or len(next_cells) != len(cells):
                break
            region.append((next_index, next_cells))
            next_position += 1
        if len(region) < 2:
            position += 1
            continue
        inferred_width = image_width or max(
            cell["bbox"]["x"] + cell["bbox"]["width"] for _index, row in region for cell in row
        )
        tolerance = max(30.0, inferred_width * 0.10)
        reference_starts = [cell["bbox"]["x"] for cell in region[0][1]]
        aligned = all(
            abs(cell["bbox"]["x"] - reference_starts[column]) <= tolerance
            for _index, row in region[1:]
            for column, cell in enumerate(row)
        )
        if not aligned:
            position += 1
            continue
        table_number = len(tables) + 1
        rows: list[list[str]] = []
        sources: list[list[dict[str, Any]]] = []
        for row_number, (source_line_index, row) in enumerate(region, 1):
            rows.append([cell["text"] for cell in row])
            sources.append(
                [
                    base_source(
                        str(path),
                        "ocr_table_cell",
                        cell["text"],
                        page=page_number,
                        table=table_number,
                        row=row_number,
                        column=column_number,
                        line=source_line_index + 1,
                        ocr=True,
                        language=language,
                        bbox=cell["bbox"],
                    )
                    for column_number, cell in enumerate(row, 1)
                ]
            )
            consumed.add(source_line_index)
        tables.append(
            NativeTable(
                title=f"{path.stem} - OCR 表格 P{page_number}-{table_number}",
                rows=rows,
                cell_sources=sources,
                kind="ocr_native_table",
                has_header=True,
            )
        )
        position = next_position
    return tables, consumed


def ocr_pdf_pages(
    pdf: Any,
    path: Path,
    page_numbers: Sequence[int],
    result: ExtractedDocument,
    language: str,
) -> list[tuple[int, dict[str, Any]]]:
    """OCR textless PDF pages.  Temporary images never leave the local machine."""
    failed: list[tuple[int, dict[str, Any]]] = []
    pages_to_process = list(page_numbers[:MAX_OCR_PAGES_PER_DOCUMENT])
    with tempfile.TemporaryDirectory(prefix="officeflow-document-ocr-") as directory:
        temp_directory = Path(directory)
        for current, page_number in enumerate(pages_to_process, 1):
            emit_progress("ocr", current - 1, len(pages_to_process), f"正在 OCR 第 {page_number} 页")
            png_path = temp_directory / f"page-{page_number:05d}.png"
            try:
                import fitz  # type: ignore

                page = pdf[page_number - 1]
                pixmap = page.get_pixmap(matrix=fitz.Matrix(OCR_RENDER_SCALE, OCR_RENDER_SCALE), alpha=False)
                pixmap.save(str(png_path))
                ocr_result = run_windows_ocr(png_path, language)
            except Exception as exc:  # noqa: BLE001 - isolate a failed OCR page
                failed.append(
                    (
                        page_number,
                        {"code": "OCR_PAGE_FAILED", "message": f"{type(exc).__name__}: {exc}"},
                    )
                )
                continue
            lines = ocr_result.get("lines", []) if ocr_result.get("success") else []
            if not isinstance(lines, list):
                lines = []
            if not lines and ocr_result.get("success") and normalize_text(ocr_result.get("text")):
                lines = [{"text": line} for line in normalize_text(ocr_result["text"]).splitlines() if line.strip()]
            valid_lines = [line for line in lines if isinstance(line, dict) and normalize_text(line.get("text"))]
            ocr_language = normalize_text(ocr_result.get("language")) or language
            inferred_tables, table_line_indexes = infer_ocr_tables(
                valid_lines,
                path,
                page_number,
                ocr_language,
                float(ocr_result.get("imageWidth", 0) or 0),
            )
            result.native_tables.extend(inferred_tables)
            added = len(table_line_indexes)
            for line_number, line in enumerate(valid_lines, 1):
                if line_number - 1 in table_line_indexes:
                    continue
                text = normalize_text(line.get("text"))
                bbox = line.get("bbox") if isinstance(line.get("bbox"), dict) else None
                result.text_units.append(
                    TextUnit(
                        text,
                        base_source(
                            str(path),
                            "ocr_line",
                            text,
                            page=page_number,
                            line=line_number,
                            ocr=True,
                            language=ocr_language,
                            bbox=bbox,
                        ),
                    )
                )
                added += 1
            if not added:
                error = ocr_result.get("error") if isinstance(ocr_result.get("error"), dict) else {}
                failed.append(
                    (
                        page_number,
                        {
                            "code": error.get("code", "OCR_NO_TEXT"),
                            "message": error.get("message", "OCR 未识别出文字。"),
                        },
                    )
                )
            emit_progress("ocr", current, len(pages_to_process), f"已 OCR 第 {page_number} 页")
    for page_number in page_numbers[MAX_OCR_PAGES_PER_DOCUMENT:]:
        failed.append(
            (
                page_number,
                {
                    "code": "OCR_PAGE_LIMIT",
                    "message": f"单文档本地 OCR 最多处理 {MAX_OCR_PAGES_PER_DOCUMENT} 页。",
                },
            )
        )
    return failed


def extract_pdf(path: Path, ocr_language: str = "zh-CN") -> ExtractedDocument:
    try:
        import fitz  # type: ignore
    except ImportError as exc:
        raise EngineError("MISSING_PYMUPDF", "处理 PDF 需要安装 PyMuPDF。") from exc

    document_result = ExtractedDocument(path=str(path))
    native_tables, table_error = pdfplumber_tables(path)
    try:
        pdf = fitz.open(str(path))
    except Exception as exc:
        raise EngineError("INVALID_PDF", "PDF 无法打开或已损坏。", {"file": str(path)}) from exc
    try:
        if len(pdf) > MAX_PDF_PAGES:
            raise EngineError(
                "PDF_TOO_MANY_PAGES",
                f"PDF 超过 {MAX_PDF_PAGES} 页限制。",
                {"file": str(path), "pages": len(pdf)},
            )
        if not native_tables:
            native_tables = pymupdf_tables(pdf, path)
        document_result.native_tables.extend(native_tables)
        duplicate_values = table_value_sets(native_tables)
        pages_without_text: list[int] = []
        for page_number, page in enumerate(pdf, 1):
            try:
                page_dict = page.get_text("dict", sort=True)
            except TypeError:
                page_dict = page.get_text("dict")
            page_line_count = 0
            for block in page_dict.get("blocks", []):
                if block.get("type", 0) != 0:
                    continue
                for line_number, line in enumerate(block.get("lines", []), 1):
                    text = normalize_text("".join(span.get("text", "") for span in line.get("spans", [])))
                    if not text:
                        continue
                    normalized = normalize_compare(text)
                    if normalized and normalized in duplicate_values.get(page_number, set()):
                        continue
                    page_line_count += 1
                    document_result.text_units.append(
                        TextUnit(
                            text,
                            base_source(
                                str(path),
                                "pdf_line",
                                text,
                                page=page_number,
                                block=int(block.get("number", 0)) + 1,
                                line=line_number,
                            ),
                        )
                    )
            if page_line_count == 0 and not any(
                int(table.cell_sources[0][0].get("page", 0)) == page_number
                for table in native_tables
                if table.cell_sources and table.cell_sources[0]
            ):
                pages_without_text.append(page_number)
        if pages_without_text:
            failed_ocr_pages = ocr_pdf_pages(pdf, path, pages_without_text, document_result, ocr_language)
            successful_count = len(pages_without_text) - len(failed_ocr_pages)
            if successful_count:
                document_result.warnings.append(
                    make_warning(
                        "LOCAL_OCR_USED",
                        f"已在本地 OCR {successful_count} 个无文本页面。",
                        file=str(path),
                    )
                )
            pages_without_text = [page_number for page_number, _error in failed_ocr_pages]
        if pages_without_text:
            failed_details = {page_number: error for page_number, error in failed_ocr_pages}
            source = base_source(
                str(path),
                "pdf_pages",
                "",
                pages=pages_without_text[:200],
                ocrErrors=[
                    {"page": page_number, **failed_details.get(page_number, {})}
                    for page_number in pages_without_text[:200]
                ],
            )
            document_result.unresolved.append(
                {
                    "reason": "OCR_REQUIRED",
                    "text": f"{len(pages_without_text)} 页本地 OCR 未能识别。",
                    "source": source,
                    "suggestedAction": "ocr",
                    "details": [
                        {"page": page_number, **failed_details.get(page_number, {})}
                        for page_number in pages_without_text[:200]
                    ],
                }
            )
            document_result.warnings.append(
                make_warning(
                    "PDF_PAGES_WITHOUT_TEXT",
                    f"有 {len(pages_without_text)} 个页面未能完成本地 OCR。",
                    file=str(path),
                    pages=pages_without_text[:200],
                )
            )
        if table_error and table_error != "pdfplumber_unavailable":
            document_result.warnings.append(
                make_warning("PDF_TABLE_FALLBACK", "PDF 表格识别已使用备用引擎。", file=str(path))
            )
    finally:
        pdf.close()
    return document_result


def extract_document(path: Path, ocr_language: str = "zh-CN") -> ExtractedDocument:
    suffix = path.suffix.casefold()
    if suffix in {".txt", ".md"}:
        return extract_plain_text(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path, ocr_language)
    raise EngineError("UNSUPPORTED_FILE_TYPE", "不支持此文件类型。", {"file": str(path)})


def unique_labels(values: Sequence[str]) -> list[str]:
    labels: list[str] = []
    counts: dict[str, int] = {}
    for index, raw in enumerate(values, 1):
        base = normalize_text(raw) or f"列{index}"
        count = counts.get(base, 0) + 1
        counts[base] = count
        labels.append(base if count == 1 else f"{base}_{count}")
    return labels


def cell_payload(value: Any, confidence: float, source_ids: Sequence[str], missing: bool = False) -> dict[str, Any]:
    return {
        "value": normalize_text(value),
        "confidence": round(max(0.0, min(1.0, confidence)), 4),
        "sourceIds": list(dict.fromkeys(source_ids)),
        **({"missing": True} if missing else {}),
    }


def make_columns(
    labels: Sequence[str], source_ids: Optional[Sequence[Sequence[str]]] = None
) -> list[dict[str, Any]]:
    columns: list[dict[str, Any]] = []
    for index, label in enumerate(unique_labels(labels), 1):
        columns.append(
            {
                "key": f"c{index}",
                "label": label,
                "sourceIds": list(source_ids[index - 1]) if source_ids and index - 1 < len(source_ids) else [],
            }
        )
    return columns


class Analyzer:
    KV_PATTERN = re.compile(r"^\s*([^:：\n]{1,40}?)\s*[:：]\s*(.*?)\s*$")
    LIST_PATTERN = re.compile(
        r"^\s*(?:(?P<number>\d{1,6})[.)、．]|(?P<cn>[一二三四五六七八九十百]{1,8})[、.)．]|[-*•▪◦])\s*(?P<text>.+?)\s*$"
    )
    MARKDOWN_SEPARATOR = re.compile(r"^\s*:?-{3,}:?\s*$")

    def __init__(self, mode: str, fields: Sequence[Any]) -> None:
        self.mode = mode
        self.fields = self._normalize_fields(fields)
        self.registry = SourceRegistry()
        self.tables: list[dict[str, Any]] = []
        self.warnings: list[dict[str, Any]] = []
        self.errors: list[dict[str, Any]] = []
        self.unresolved: list[dict[str, Any]] = []
        self.total_units = 0
        self.structured_units = 0
        self.preserved_units = 0
        self.files_analyzed = 0

    @staticmethod
    def _normalize_fields(fields: Sequence[Any]) -> list[dict[str, Any]]:
        normalized: list[dict[str, Any]] = []
        seen: set[str] = set()
        for index, item in enumerate(fields):
            if isinstance(item, str):
                label = normalize_text(item)
                key = f"field{index + 1}"
                aliases: list[str] = []
            elif isinstance(item, dict):
                label = normalize_text(item.get("label") or item.get("key"))
                key = normalize_text(item.get("key")) or f"field{index + 1}"
                raw_aliases = item.get("aliases", [])
                aliases = [normalize_text(value) for value in raw_aliases if isinstance(value, str)] if isinstance(raw_aliases, list) else []
            else:
                continue
            if not label:
                continue
            canonical = normalize_key(key)
            if canonical in seen:
                continue
            seen.add(canonical)
            normalized.append({"key": key, "label": label, "aliases": list(dict.fromkeys([label, *aliases]))})
        return normalized

    def add_source(self, source: dict[str, Any]) -> str:
        return self.registry.add(source)

    def add_unresolved(
        self,
        reason: str,
        text: str,
        source_ids: Sequence[str],
        suggested_action: str = "review",
    ) -> None:
        self.unresolved.append(
            {
                "id": f"unresolved-{len(self.unresolved) + 1:06d}",
                "reason": reason,
                "text": normalize_text(text),
                "sourceIds": list(dict.fromkeys(source_ids)),
                "suggestedAction": suggested_action,
            }
        )

    def analyze_document(self, document: ExtractedDocument) -> None:
        self.files_analyzed += 1
        self.warnings.extend(document.warnings)
        for item in document.unresolved:
            source_id = self.add_source(item["source"])
            self.add_unresolved(item["reason"], item["text"], [source_id], item.get("suggestedAction", "review"))

        native_units = sum(max(1, len(table.rows) - (1 if table.has_header and len(table.rows) > 1 else 0)) for table in document.native_tables)
        text_nonempty = sum(1 for unit in document.text_units if unit.text.strip())
        self.total_units += native_units + text_nonempty + len(document.unresolved)
        self.preserved_units += native_units + len(document.unresolved)
        self.structured_units += native_units

        for table in document.native_tables:
            rendered = self.render_native_table(table)
            if rendered:
                self.tables.append(rendered)

        if not document.text_units:
            return
        consumed: set[int] = set()
        if self.mode == "ledger":
            self.render_ledger(document.text_units, consumed, mark_unresolved=False)
            return

        for parsed_table, indexes in self.find_delimited_tables(document.text_units):
            self.tables.append(self.render_native_table(parsed_table))
            consumed.update(indexes)
            count = sum(1 for index in indexes if document.text_units[index].text.strip())
            self.structured_units += count
            self.preserved_units += count

        if self.mode == "native":
            self.render_unresolved_only(document.text_units, consumed, "NO_NATIVE_TABLE")
            return

        self.render_key_value_tables(document.text_units, consumed)
        if self.fields:
            self.render_typed_field_rows(document.text_units, consumed)
        self.render_lists(document.text_units, consumed)
        self.render_ledger(document.text_units, consumed, mark_unresolved=True)

    def render_native_table(self, table: NativeTable) -> dict[str, Any]:
        table_id = f"table-{len(self.tables) + 1:04d}"
        width = max((len(row) for row in table.rows), default=0)
        if width == 0:
            return {}
        use_header = bool(table.has_header and len(table.rows) >= 2)
        header = table.rows[0] if use_header else [f"列{index}" for index in range(1, width + 1)]
        header_sources = table.cell_sources[0] if use_header and table.cell_sources else []
        header_ids = [[self.add_source(source)] for source in header_sources]
        columns = make_columns([header[index] if index < len(header) else "" for index in range(width)], header_ids)
        data_start = 1 if use_header else 0
        rendered_rows: list[dict[str, Any]] = []
        for row_index in range(data_start, len(table.rows)):
            values = table.rows[row_index]
            sources = table.cell_sources[row_index] if row_index < len(table.cell_sources) else []
            cells: dict[str, dict[str, Any]] = {}
            row_source_ids: list[str] = []
            for column_index, column in enumerate(columns):
                value = values[column_index] if column_index < len(values) else ""
                if column_index < len(sources):
                    source_id = self.add_source(sources[column_index])
                elif sources:
                    source_id = self.add_source(sources[-1])
                else:
                    fallback = base_source(".", "unknown", value, row=row_index + 1, column=column_index + 1)
                    source_id = self.add_source(fallback)
                row_source_ids.append(source_id)
                cells[column["key"]] = cell_payload(value, 0.99, [source_id], missing=not bool(value))
            rendered_rows.append(
                {
                    "id": f"{table_id}-row-{row_index - data_start + 1}",
                    "cells": cells,
                    "sourceIds": list(dict.fromkeys(row_source_ids)),
                    "status": "ready",
                }
            )
        return {
            "id": table_id,
            "title": table.title,
            "kind": table.kind,
            "columns": columns,
            "rows": rendered_rows,
            "rowCount": len(rendered_rows),
        }

    @staticmethod
    def parse_delimited_line(text: str, delimiter: str) -> list[str]:
        stripped = text.strip()
        if delimiter == "|":
            stripped = stripped.strip("|")
        try:
            return [normalize_text(value) for value in next(csv.reader([stripped], delimiter=delimiter, skipinitialspace=True))]
        except (csv.Error, StopIteration):
            return []

    def detect_delimiter(self, lines: Sequence[str]) -> Optional[str]:
        for delimiter in ("|", "\t", ";", ","):
            parsed = [self.parse_delimited_line(line, delimiter) for line in lines]
            widths = [len(row) for row in parsed]
            if not widths or min(widths) < 2 or len(set(widths)) != 1:
                continue
            if delimiter == "," and len(lines) < 3:
                continue
            if delimiter == "|" and not all("|" in line for line in lines):
                continue
            return delimiter
        return None

    def find_delimited_tables(self, units: Sequence[TextUnit]) -> Iterable[tuple[NativeTable, set[int]]]:
        start = 0
        while start < len(units):
            while start < len(units) and not units[start].text.strip():
                start += 1
            end = start
            while end < len(units) and units[end].text.strip():
                end += 1
            if end - start >= 2:
                lines = [units[index].text for index in range(start, end)]
                delimiter = self.detect_delimiter(lines)
                if delimiter:
                    parsed = [self.parse_delimited_line(line, delimiter) for line in lines]
                    separator_index: Optional[int] = None
                    if delimiter == "|" and len(parsed) >= 2 and all(
                        self.MARKDOWN_SEPARATOR.fullmatch(value) for value in parsed[1]
                    ):
                        separator_index = 1
                    rows: list[list[str]] = []
                    sources: list[list[dict[str, Any]]] = []
                    indexes: set[int] = set()
                    for offset, values in enumerate(parsed):
                        index = start + offset
                        indexes.add(index)
                        if offset == separator_index:
                            continue
                        rows.append(values)
                        sources.append([{**units[index].source, "column": column + 1, "text": value} for column, value in enumerate(values)])
                    if len(rows) >= 2:
                        yield (
                            NativeTable(
                                title=f"{Path(units[start].source['file']).stem} - 分隔表",
                                rows=rows,
                                cell_sources=sources,
                                kind="markdown_table" if separator_index is not None else "delimited_table",
                                has_header=True,
                            ),
                            indexes,
                        )
            start = max(end + 1, start + 1)

    def field_for_label(self, label: str) -> Optional[dict[str, Any]]:
        canonical = normalize_key(label)
        for field_definition in self.fields:
            if any(normalize_key(alias) == canonical for alias in field_definition["aliases"]):
                return field_definition
        return None

    def collect_key_value_records(
        self, units: Sequence[TextUnit], consumed: set[int]
    ) -> list[list[tuple[str, str, int]]]:
        records: list[list[tuple[str, str, int]]] = []
        current: list[tuple[str, str, int]] = []
        current_keys: set[str] = set()
        for index, unit in enumerate(units):
            if index in consumed:
                continue
            text = unit.text.strip()
            if not text:
                if current:
                    records.append(current)
                    current = []
                    current_keys = set()
                continue
            match = self.KV_PATTERN.fullmatch(text)
            if not match:
                if current:
                    records.append(current)
                    current = []
                    current_keys = set()
                continue
            label = normalize_text(match.group(1))
            value = normalize_text(match.group(2))
            canonical = normalize_key(label)
            if canonical in current_keys and current:
                records.append(current)
                current = []
                current_keys = set()
            current.append((label, value, index))
            current_keys.add(canonical)
        if current:
            records.append(current)
        return [record for record in records if len(record) >= 2]

    def render_key_value_tables(self, units: Sequence[TextUnit], consumed: set[int]) -> None:
        records = self.collect_key_value_records(units, consumed)
        if not records:
            return
        groups = [records] if self.fields else self.cluster_key_value_records(records)
        for group_number, group in enumerate(groups, 1):
            self.render_key_value_group(units, consumed, group, group_number, len(groups))

    @staticmethod
    def cluster_key_value_records(
        records: Sequence[list[tuple[str, str, int]]]
    ) -> list[list[list[tuple[str, str, int]]]]:
        """Keep unrelated key/value schemas out of the same sparse table."""
        groups: list[dict[str, Any]] = []
        for record in records:
            schema = {normalize_key(label) for label, _value, _index in record}
            best_group: Optional[dict[str, Any]] = None
            best_score = 0.0
            for group in groups:
                known_schema: set[str] = group["schema"]
                intersection = len(schema & known_schema)
                union = len(schema | known_schema)
                score = intersection / union if union else 0.0
                if intersection >= 2 and score >= best_score:
                    best_group = group
                    best_score = score
            if best_group is None:
                groups.append({"schema": set(schema), "records": [record]})
            else:
                best_group["records"].append(record)
                best_group["schema"].update(schema)
        return [group["records"] for group in groups]

    def render_key_value_group(
        self,
        units: Sequence[TextUnit],
        consumed: set[int],
        records: Sequence[list[tuple[str, str, int]]],
        group_number: int,
        group_count: int,
    ) -> None:
        # Auto mode is intentionally conservative: a two-field singleton often
        # comes from a heading such as "Page: 1 / Duration: 30s", not a record.
        # Keep it in the review ledger unless the user explicitly chose fields.
        if not self.fields and len(records) == 1 and len(records[0]) < 3:
            return
        table_id = f"table-{len(self.tables) + 1:04d}"
        if self.fields:
            definitions = self.fields
        else:
            labels: list[str] = []
            seen: set[str] = set()
            for record in records:
                for label, _value, _index in record:
                    canonical = normalize_key(label)
                    if canonical not in seen:
                        seen.add(canonical)
                        labels.append(label)
            definitions = [{"key": f"field{index + 1}", "label": label, "aliases": [label]} for index, label in enumerate(labels)]

        columns = [
            {"key": f"c{index + 1}", "label": definition["label"], "fieldKey": definition["key"], "sourceIds": []}
            for index, definition in enumerate(definitions)
        ]
        rows: list[dict[str, Any]] = []
        used_count = 0
        for record_number, record in enumerate(records, 1):
            record_values: dict[str, tuple[str, int]] = {}
            for label, value, unit_index in record:
                definition = self.field_for_label(label) if self.fields else next(
                    (item for item in definitions if normalize_key(item["label"]) == normalize_key(label)),
                    None,
                )
                if definition:
                    record_values[definition["key"]] = (value, unit_index)
            if not record_values:
                continue
            row_source_ids = [self.add_source(units[index].source) for _label, _value, index in record]
            fallback_source_id = row_source_ids[0]
            cells: dict[str, dict[str, Any]] = {}
            missing = False
            for column, definition in zip(columns, definitions):
                found = record_values.get(definition["key"])
                if found:
                    value, unit_index = found
                    source_id = self.add_source(units[unit_index].source)
                    cells[column["key"]] = cell_payload(value, 0.98, [source_id])
                else:
                    missing = True
                    cells[column["key"]] = cell_payload("", 0.0, [fallback_source_id], missing=True)
            rows.append(
                {
                    "id": f"{table_id}-row-{record_number}",
                    "cells": cells,
                    "sourceIds": list(dict.fromkeys(row_source_ids)),
                    "status": "needs_review" if missing else "ready",
                }
            )
            record_indexes = {index for _label, _value, index in record}
            consumed.update(record_indexes)
            used_count += len(record_indexes)
            if missing:
                self.add_unresolved("MISSING_FIELDS", "存在无法本地提取的字段。", row_source_ids, "review_or_ai")
        if rows:
            first_index = records[0][0][2]
            file_stem = Path(units[first_index].source.get("file", "文档")).stem
            title_suffix = f" {group_number}" if group_count > 1 else ""
            self.tables.append(
                {
                    "id": table_id,
                    "title": (
                        f"{file_stem} - 自定义字段提取"
                        if self.fields
                        else f"{file_stem} - 重复键值记录{title_suffix}"
                    ),
                    "kind": "custom_fields" if self.fields else "key_value_records",
                    "columns": columns,
                    "rows": rows,
                    "rowCount": len(rows),
                }
            )
            self.structured_units += used_count
            self.preserved_units += used_count

    @staticmethod
    def known_field_pattern(label: str) -> Optional[re.Pattern[str]]:
        canonical = normalize_key(label)
        if any(token in canonical for token in ("电话", "手机", "mobile", "phone")):
            return re.compile(r"(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)")
        if any(token in canonical for token in ("邮箱", "邮件", "email")):
            return re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
        if any(token in canonical for token in ("日期", "时间", "截止", "date")):
            return re.compile(r"(?:20\d{2}[-/.年]\d{1,2}[-/.月]\d{1,2}日?|\d{1,2}月\d{1,2}日)")
        if any(token in canonical for token in ("金额", "价格", "费用", "合计", "amount", "price")):
            return re.compile(r"(?:¥|￥)?\s*\d[\d,]*(?:\.\d{1,2})?\s*(?:元|万元|人民币)?")
        if any(token in canonical for token in ("身份证", "证件号")):
            return re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)")
        return None

    def extract_field_from_text(self, definition: dict[str, Any], text: str) -> Optional[str]:
        aliases = sorted(definition["aliases"], key=len, reverse=True)
        alias_pattern = "|".join(re.escape(alias) for alias in aliases if alias)
        if alias_pattern:
            match = re.search(
                rf"(?:^|[;；,，。\s])(?:{alias_pattern})\s*[:：]\s*([^;；,，。\n]+)",
                text,
                flags=re.IGNORECASE,
            )
            if match:
                return normalize_text(match.group(1))
        known_pattern = self.known_field_pattern(definition["label"])
        if known_pattern:
            match = known_pattern.search(text)
            if match:
                return normalize_text(match.group(0))
        return None

    def render_typed_field_rows(self, units: Sequence[TextUnit], consumed: set[int]) -> None:
        if not self.fields:
            return
        table_id = f"table-{len(self.tables) + 1:04d}"
        columns = [
            {"key": f"c{index + 1}", "label": definition["label"], "fieldKey": definition["key"], "sourceIds": []}
            for index, definition in enumerate(self.fields)
        ]
        rows: list[dict[str, Any]] = []
        for index, unit in enumerate(units):
            if index in consumed or not unit.text.strip():
                continue
            found_values = [self.extract_field_from_text(definition, unit.text) for definition in self.fields]
            if not any(value is not None for value in found_values):
                continue
            source_id = self.add_source(unit.source)
            cells: dict[str, dict[str, Any]] = {}
            missing = False
            for column, value in zip(columns, found_values):
                if value is None:
                    missing = True
                    cells[column["key"]] = cell_payload("", 0.0, [source_id], missing=True)
                else:
                    cells[column["key"]] = cell_payload(value, 0.90, [source_id])
            rows.append(
                {
                    "id": f"{table_id}-row-{len(rows) + 1}",
                    "cells": cells,
                    "sourceIds": [source_id],
                    "status": "needs_review" if missing else "ready",
                }
            )
            consumed.add(index)
            self.structured_units += 1
            self.preserved_units += 1
            if missing:
                self.add_unresolved("MISSING_FIELDS", unit.text, [source_id], "review_or_ai")
        if rows:
            file_stem = Path(units[next(iter(consumed))].source.get("file", "文档")).stem if consumed else "文档"
            self.tables.append(
                {
                    "id": table_id,
                    "title": f"{file_stem} - 自定义字段提取（文本）",
                    "kind": "custom_fields_text",
                    "columns": columns,
                    "rows": rows,
                    "rowCount": len(rows),
                }
            )

    def render_lists(self, units: Sequence[TextUnit], consumed: set[int]) -> None:
        table_id = f"table-{len(self.tables) + 1:04d}"
        rows: list[dict[str, Any]] = []
        indexes: list[int] = []
        for index, unit in enumerate(units):
            if index in consumed:
                continue
            match = self.LIST_PATTERN.fullmatch(unit.text)
            if not match:
                continue
            source_id = self.add_source(unit.source)
            ordinal = match.group("number") or match.group("cn") or str(len(rows) + 1)
            rows.append(
                {
                    "id": f"{table_id}-row-{len(rows) + 1}",
                    "cells": {
                        "c1": cell_payload(ordinal, 0.97, [source_id]),
                        "c2": cell_payload(match.group("text"), 0.97, [source_id]),
                    },
                    "sourceIds": [source_id],
                    "status": "ready",
                }
            )
            indexes.append(index)
        if rows:
            file_stem = Path(units[indexes[0]].source.get("file", "文档")).stem
            self.tables.append(
                {
                    "id": table_id,
                    "title": f"{file_stem} - 列表台账",
                    "kind": "list_ledger",
                    "columns": make_columns(["序号", "内容"]),
                    "rows": rows,
                    "rowCount": len(rows),
                }
            )
            consumed.update(indexes)
            self.structured_units += len(indexes)
            self.preserved_units += len(indexes)

    def render_ledger(self, units: Sequence[TextUnit], consumed: set[int], mark_unresolved: bool) -> None:
        table_id = f"table-{len(self.tables) + 1:04d}"
        rows: list[dict[str, Any]] = []
        for index, unit in enumerate(units):
            if index in consumed or not unit.text.strip():
                continue
            source_id = self.add_source(unit.source)
            rows.append(
                {
                    "id": f"{table_id}-row-{len(rows) + 1}",
                    "cells": {
                        "c1": cell_payload(str(len(rows) + 1), 1.0, [source_id]),
                        "c2": cell_payload(unit.text, 0.60 if mark_unresolved else 1.0, [source_id]),
                    },
                    "sourceIds": [source_id],
                    "status": "needs_review" if mark_unresolved else "ready",
                }
            )
            consumed.add(index)
            self.preserved_units += 1
            if not mark_unresolved:
                self.structured_units += 1
            else:
                self.add_unresolved("LOW_STRUCTURE_TEXT", unit.text, [source_id], "review_or_ai")
        if rows:
            first_source_id = rows[0]["sourceIds"][0]
            first_source = self.registry.as_dict().get(first_source_id, {})
            file_stem = Path(first_source.get("file", "文档")).stem
            self.tables.append(
                {
                    "id": table_id,
                    "title": f"{file_stem} - 段落台账",
                    "kind": "paragraph_ledger",
                    "columns": make_columns(["序号", "内容"]),
                    "rows": rows,
                    "rowCount": len(rows),
                }
            )

    def render_unresolved_only(self, units: Sequence[TextUnit], consumed: set[int], reason: str) -> None:
        for index, unit in enumerate(units):
            if index in consumed or not unit.text.strip():
                continue
            source_id = self.add_source(unit.source)
            self.preserved_units += 1
            self.add_unresolved(reason, unit.text, [source_id], "review")

    def result(self, files_requested: int, elapsed_ms: float) -> dict[str, Any]:
        source_coverage = self.preserved_units / self.total_units if self.total_units else 1.0
        structural_coverage = self.structured_units / self.total_units if self.total_units else 1.0
        stats = {
            "filesRequested": files_requested,
            "filesAnalyzed": self.files_analyzed,
            "tableCount": len(self.tables),
            "totalRows": sum(len(table.get("rows", [])) for table in self.tables),
            "totalSourceUnits": self.total_units,
            "preservedSourceUnits": self.preserved_units,
            "structuredSourceUnits": self.structured_units,
            "sourceCoverage": round(min(1.0, source_coverage), 4),
            "structuralCoverage": round(min(1.0, structural_coverage), 4),
            "unresolvedCount": len(self.unresolved),
            "needsAi": bool(self.unresolved),
            "elapsedMs": round(elapsed_ms, 3),
        }
        result: dict[str, Any] = {
            "version": 1,
            "engineVersion": ENGINE_VERSION,
            "success": not self.errors,
            "generatedAt": utc_now(),
            "tables": self.tables,
            "warnings": self.warnings,
            "errors": self.errors,
            "unresolved": self.unresolved,
            "stats": stats,
            "sources": self.registry.as_dict(),
        }
        if len(self.tables) == 1:
            main_table = self.tables[0]
            result["columns"] = main_table["columns"]
            result["rows"] = main_table["rows"]
            result["totalRows"] = len(main_table["rows"])
        else:
            result["columns"] = []
            result["rows"] = []
            result["totalRows"] = stats["totalRows"]
        return result


def validate_analyze_request(request: dict[str, Any]) -> tuple[list[Any], str, list[Any], str]:
    files = request.get("files")
    if not isinstance(files, list) or not files:
        raise EngineError("FILES_REQUIRED", "files 必须是至少包含一个路径的数组。")
    if len(files) > MAX_FILES:
        raise EngineError("TOO_MANY_FILES", f"单次最多处理 {MAX_FILES} 个文件。")
    mode = request.get("mode", "auto")
    if mode not in VALID_MODES:
        raise EngineError("INVALID_MODE", f"mode 必须是 {', '.join(sorted(VALID_MODES))} 之一。")
    fields = request.get("fields", [])
    if not isinstance(fields, list):
        raise EngineError("INVALID_FIELDS", "fields 必须是数组。")
    if mode == "fields" and not fields:
        raise EngineError("FIELDS_REQUIRED", "fields 模式必须提供至少一个字段。")
    if len(fields) > 100:
        raise EngineError("TOO_MANY_FIELDS", "自定义字段最多 100 个。")
    ocr_language = normalize_text(request.get("ocrLanguage", "zh-CN")) or "zh-CN"
    if not re.fullmatch(r"[A-Za-z]{2,3}(?:-[A-Za-z0-9]{2,8})?", ocr_language):
        raise EngineError("INVALID_OCR_LANGUAGE", "ocrLanguage 必须是有效语言标签，例如 zh-CN。")
    return files, mode, fields, ocr_language


def analyze_request(request: dict[str, Any]) -> dict[str, Any]:
    started = time.perf_counter()
    files, mode, fields, ocr_language = validate_analyze_request(request)
    analyzer = Analyzer(mode, fields)
    for index, raw_path in enumerate(files, 1):
        display_name = Path(raw_path).name if isinstance(raw_path, str) else f"文件 {index}"
        emit_progress("extract", index - 1, len(files), f"正在读取 {display_name}")
        try:
            path = validate_input_path(raw_path)
            document = extract_document(path, ocr_language)
            analyzer.analyze_document(document)
        except EngineError as exc:
            analyzer.errors.append({**exc.as_dict(), "file": str(raw_path)})
        except Exception as exc:  # noqa: BLE001 - isolate failures to one input file
            analyzer.errors.append(
                {
                    "code": "FILE_ANALYSIS_FAILED",
                    "message": "文档分析失败。",
                    "file": str(raw_path),
                    "detail": f"{type(exc).__name__}: {exc}",
                }
            )
        emit_progress("analyze", index, len(files), f"已分析 {display_name}")
    return analyzer.result(len(files), (time.perf_counter() - started) * 1000)


def load_analysis(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        analysis = value
    elif isinstance(value, str):
        analysis = read_json_file(value)
    else:
        raise EngineError("INVALID_ANALYSIS", "analysis 必须是分析结果对象或 JSON 文件路径。")
    if not isinstance(analysis.get("tables"), list):
        raise EngineError("INVALID_ANALYSIS", "analysis.tables 必须是数组。")
    return analysis


def output_path_for_format(
    raw_path: Any, file_format: str
) -> tuple[Path, Optional[dict[str, Any]]]:
    if not isinstance(raw_path, str) or not raw_path.strip():
        raise EngineError("OUTPUT_PATH_REQUIRED", "outputPath 必须是非空路径。")
    target = Path(raw_path).expanduser().resolve()
    expected_suffix = ".xlsx" if file_format == "xlsx" else ".csv"
    warning = None
    if target.suffix.casefold() != expected_suffix:
        target = target.with_suffix(expected_suffix)
        warning = make_warning("OUTPUT_EXTENSION_ADJUSTED", f"输出扩展名已调整为 {expected_suffix}。", outputPath=str(target))
    target.parent.mkdir(parents=True, exist_ok=True)
    return target, warning


def source_ids_for_row(row: dict[str, Any]) -> list[str]:
    ids = list(row.get("sourceIds", [])) if isinstance(row.get("sourceIds"), list) else []
    if ids:
        return list(dict.fromkeys(str(value) for value in ids))
    for cell in (row.get("cells") or {}).values():
        if isinstance(cell, dict) and isinstance(cell.get("sourceIds"), list):
            ids.extend(str(value) for value in cell["sourceIds"])
    return list(dict.fromkeys(ids))


def export_csv(analysis: dict[str, Any], target: Path) -> None:
    tables = analysis.get("tables", [])
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        for table_index, table in enumerate(tables):
            if table_index:
                writer.writerow([])
            if len(tables) > 1:
                writer.writerow(["表格", safe_cell_text(table.get("title", f"表格 {table_index + 1}"))])
            columns = table.get("columns", [])
            writer.writerow([safe_cell_text(column.get("label", "")) for column in columns] + ["来源编号"])
            for row in table.get("rows", []):
                cells = row.get("cells", {})
                values = [safe_cell_text((cells.get(column.get("key")) or {}).get("value", "")) for column in columns]
                writer.writerow(values + [";".join(source_ids_for_row(row))])


def excel_sheet_title(value: str, fallback: str, existing: set[str]) -> str:
    clean = re.sub(r"[\\/*?:\[\]]", "_", normalize_text(value))[:31] or fallback
    candidate = clean
    suffix = 2
    while candidate.casefold() in existing:
        addition = f"_{suffix}"
        candidate = clean[: 31 - len(addition)] + addition
        suffix += 1
    existing.add(candidate.casefold())
    return candidate


def style_worksheet(worksheet: Any, max_column: int, max_row: int) -> None:
    from openpyxl.styles import (  # type: ignore
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )

    header_fill = PatternFill("solid", fgColor="0B78D1")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2EA")
    border = Border(bottom=thin)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2, max_row=max_row, max_col=max_column):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    worksheet.freeze_panes = "A2"
    if max_row >= 1 and max_column >= 1:
        worksheet.auto_filter.ref = f"A1:{worksheet.cell(max_row, max_column).coordinate}"
    for column_index in range(1, max_column + 1):
        length = 10
        for row_index in range(1, min(max_row, 500) + 1):
            value = worksheet.cell(row_index, column_index).value
            if value is not None:
                parts = str(value).split("\n")
                display_length = max(
                    sum(2 if unicodedata.east_asian_width(char) in {"W", "F", "A"} else 1 for char in part)
                    for part in parts
                )
                length = max(length, min(60, display_length + 2))
        worksheet.column_dimensions[worksheet.cell(1, column_index).column_letter].width = min(60, length)
    worksheet.sheet_view.showGridLines = False


def export_xlsx(analysis: dict[str, Any], target: Path) -> None:
    try:
        from openpyxl import Workbook  # type: ignore
        from openpyxl.comments import Comment  # type: ignore
    except ImportError as exc:
        raise EngineError("MISSING_OPENPYXL", "导出 XLSX 需要安装 openpyxl。") from exc

    workbook = Workbook()
    workbook.remove(workbook.active)
    # Reserve audit sheet names so a document table cannot displace them.
    existing_titles: set[str] = {"待确认".casefold(), "来源索引".casefold()}
    tables = analysis.get("tables", [])
    if not tables:
        worksheet = workbook.create_sheet("数据")
        worksheet.append(["提示"])
        worksheet.append(["没有可导出的结构化数据"])
        style_worksheet(worksheet, 1, 2)
        existing_titles.add("数据".casefold())
    for table_index, table in enumerate(tables):
        title = "数据" if table_index == 0 else excel_sheet_title(table.get("title", ""), f"数据_{table_index + 1}", existing_titles)
        if table_index == 0:
            existing_titles.add(title.casefold())
        worksheet = workbook.create_sheet(title)
        columns = table.get("columns", [])
        worksheet.append([safe_cell_text(column.get("label", "")) for column in columns] + ["来源编号"])
        for row in table.get("rows", []):
            cells = row.get("cells", {})
            values: list[str] = []
            comments: list[str] = []
            for column in columns:
                cell = cells.get(column.get("key"), {})
                values.append(safe_cell_text(cell.get("value", "")))
                comments.append("、".join(str(value) for value in cell.get("sourceIds", [])))
            row_source_ids = source_ids_for_row(row)
            worksheet.append(values + [";".join(row_source_ids)])
            excel_row = worksheet.max_row
            for column_index, source_comment in enumerate(comments, 1):
                if source_comment:
                    worksheet.cell(excel_row, column_index).comment = Comment(f"来源：{source_comment}", "OfficeFlow")
        style_worksheet(worksheet, max(1, len(columns) + 1), worksheet.max_row)

    pending = workbook.create_sheet("待确认")
    pending.append(["编号", "原因", "内容", "来源编号", "建议操作"])
    for item in analysis.get("unresolved", []):
        pending.append(
            [
                safe_cell_text(item.get("id", "")),
                safe_cell_text(item.get("reason", "")),
                safe_cell_text(item.get("text", "")),
                safe_cell_text(";".join(str(value) for value in item.get("sourceIds", []))),
                safe_cell_text(item.get("suggestedAction", "review")),
            ]
        )
    for warning in analysis.get("warnings", []):
        pending.append(
            [
                "",
                safe_cell_text(warning.get("code", "WARNING")),
                safe_cell_text(warning.get("message", "")),
                "",
                "review",
            ]
        )
    style_worksheet(pending, 5, pending.max_row)

    source_sheet = workbook.create_sheet("来源索引")
    source_sheet.append(["来源编号", "文件", "类型", "页码", "段落", "行号", "表格", "行", "列", "原文"])
    sources = analysis.get("sources", {})
    if isinstance(sources, dict):
        source_items = sources.items()
    else:
        source_items = ((str(item.get("id", "")), item) for item in sources if isinstance(item, dict))
    for source_id, source in source_items:
        source_sheet.append(
            [
                safe_cell_text(source_id),
                safe_cell_text(source.get("file", "")),
                safe_cell_text(source.get("type", "")),
                source.get("page", ""),
                source.get("paragraph", ""),
                source.get("lineStart", source.get("line", "")),
                source.get("table", ""),
                source.get("row", ""),
                source.get("column", ""),
                safe_cell_text(source.get("text", "")),
            ]
        )
    style_worksheet(source_sheet, 10, source_sheet.max_row)

    workbook.properties.creator = "OfficeFlow"
    workbook.properties.title = "OfficeFlow 文档整理结果"
    workbook.properties.description = "本地文档转表格结果，包含来源追溯与待确认项。"
    workbook.save(target)


def export_request(request: dict[str, Any]) -> dict[str, Any]:
    started = time.perf_counter()
    analysis = load_analysis(request.get("analysis"))
    file_format = normalize_text(request.get("format", "xlsx")).casefold()
    if file_format not in {"xlsx", "csv"}:
        raise EngineError("INVALID_EXPORT_FORMAT", "format 必须是 xlsx 或 csv。")
    target, extension_warning = output_path_for_format(request.get("outputPath"), file_format)
    warnings = [extension_warning] if extension_warning else []
    emit_progress("export", 0, 1, f"正在生成 {target.name}")
    temporary = target.with_name(f".{target.name}.{os.getpid()}.tmp{target.suffix}")
    try:
        if file_format == "xlsx":
            export_xlsx(analysis, temporary)
        else:
            export_csv(analysis, temporary)
        os.replace(temporary, target)
    except Exception:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass
        raise
    emit_progress("export", 1, 1, f"已生成 {target.name}")
    return {
        "version": 1,
        "engineVersion": ENGINE_VERSION,
        "success": True,
        "generatedAt": utc_now(),
        "outputs": [str(target)],
        "warnings": warnings,
        "errors": [],
        "stats": {
            "format": file_format,
            "tableCount": len(analysis.get("tables", [])),
            "rowCount": sum(len(table.get("rows", [])) for table in analysis.get("tables", [])),
            "elapsedMs": round((time.perf_counter() - started) * 1000, 3),
        },
    }


def failure_result(error: EngineError) -> dict[str, Any]:
    return {
        "version": 1,
        "engineVersion": ENGINE_VERSION,
        "success": False,
        "generatedAt": utc_now(),
        "tables": [],
        "columns": [],
        "rows": [],
        "totalRows": 0,
        "warnings": [],
        "errors": [error.as_dict()],
        "unresolved": [],
        "stats": {"needsAi": False},
        "sources": {},
    }


def main(argv: Sequence[str]) -> int:
    if len(argv) != 4 or argv[1] not in {"analyze", "export"}:
        print("Usage: document_table.py <analyze|export> <request.json> <result.json>", file=sys.stderr)
        return 2
    action, request_path, result_path = argv[1], argv[2], argv[3]
    try:
        request = read_json_file(request_path)
        result = analyze_request(request) if action == "analyze" else export_request(request)
        write_json_atomic(result_path, result)
        return 0 if result.get("success") else 1
    except EngineError as exc:
        write_json_atomic(result_path, failure_result(exc))
        print(f"{exc.code}: {exc.message}", file=sys.stderr)
        return 1
    except Exception as exc:  # noqa: BLE001 - CLI boundary must always write a result file
        error = EngineError("UNEXPECTED_ERROR", "处理过程中发生未预期错误。", {"detail": f"{type(exc).__name__}: {exc}"})
        write_json_atomic(result_path, failure_result(error))
        print(f"UNEXPECTED_ERROR: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
